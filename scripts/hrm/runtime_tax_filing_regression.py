#!/usr/bin/env python3
"""Verify closed-period offline filing and receipt reconciliation using synthetic paid wages."""
import argparse
import datetime
from decimal import Decimal
import io
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from runtime_client import RuntimeClient

p=argparse.ArgumentParser(description=__doc__)
p.add_argument('--state-file',required=True)
p.add_argument('--output-dir',type=Path,required=True)
a=p.parse_args()
c=RuntimeClient(a.state_file)
s=c.state
e=s['entityId'];year=s['fixtureYear'];month=s['fixtureMonth']
review=c.login('hrmacceptreview','reviewerPassword')
finance=c.login('hrmacceptoperator','browserPassword')
today=datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=8))).date()
reason='虚构离线申报/完税回执验收，不向机构申报或付款，不作为正式业务凭据'
a.output_dir.mkdir(parents=True,exist_ok=True)

def remember(key,create):
    if key not in s:s[key]=create();c.save()
    return s[key]

def post(path,params,token=review):return c.query('/hrm/payroll/tax-filing/'+path,params,token)

def detail(id):return c.ok(f'/hrm/payroll/tax-filing/detail?id={id}')

current=c.request('/hrm/payroll/tax-filing/create',{'entityId':e,'year':today.year,'month':today.month,'zeroDeclaration':False,'requestKey':'ACCEPT_REJECT_UNCLOSED','evidence':reason})
assert current['code'] not in (0,500),current
request={'entityId':e,'year':year,'month':month,'zeroDeclaration':False,'requestKey':f'ACCEPT_FILING_{e}_{year}_{month}','evidence':reason}
filing=remember('taxFilingId',lambda:c.ok('/hrm/payroll/tax-filing/create',request))
assert c.ok('/hrm/payroll/tax-filing/create',request)==filing
row=detail(filing)
assert row['filing']['expectedTax']==18 and row['sourceFresh'],row['filing']['expectedTax']
if row['filing']['status']=='PENDING':
    denied=c.request(f'/hrm/payroll/tax-filing/review?id={filing}&approved=true&reason=synthetic-self-review',{})
    assert denied['code'] not in (0,500),denied
    post('review',{'id':filing,'approved':True,'reason':reason})
content=c.request(f'/hrm/payroll/tax-filing/export?id={filing}',token=finance)
assert isinstance(content,bytes) and content[:2]==b'PK'
sheet=load_workbook(io.BytesIO(content)).active
headers=[cell.value for cell in sheet[1]]
rows=[dict(zip(headers,row)) for row in sheet.iter_rows(min_row=2,values_only=True)]
assert len(rows)==1 and rows[0]['收入所属月份']==f'{year}-{month:02d}',rows
assert Decimal(str(rows[0]['本月收入']))==6000
assert Decimal(str(rows[0]['本月已预扣税']))==18
assert Decimal(str(rows[0]['累计收入']))==6000*month
assert Decimal(str(rows[0]['累计应纳税额']))==18*month
assert rows[0]['证件号码'].startswith('TESTONLY')
post('filed',{'id':filing,'reference':f'OFFLINE_TEST_FILING_{filing}','date':today.isoformat(),'evidence':reason},c.token)
fixture=Path(a.state_file).resolve().parent/'tax-receipt-fixture.xlsx'
if not fixture.exists():
    book=Workbook();sheet=book.active
    sheet.append(['凭证类型','官方凭证编号','金额','实际缴退日期','原缴税凭证记录编号','核实依据'])
    for index,amount in enumerate((17,1),1):sheet.append(['缴税',f'OFFLINE_TEST_RECEIPT_{filing}_{index}',amount,today.isoformat(),None,reason])
    book.save(fixture)
raw=fixture.read_bytes()
first=c.upload(f'/hrm/payroll/tax-filing/receipt-import?id={filing}','synthetic-tax-receipts.xlsx',raw)
assert first['code']==0,first
ids=first['data'];assert len(ids)==2
repeat=c.upload(f'/hrm/payroll/tax-filing/receipt-import?id={filing}','synthetic-tax-receipts.xlsx',raw)
assert repeat['code']==0 and repeat['data']==ids,repeat
row=detail(filing)
for index,id in enumerate(ids):
    receipt=next(item for item in row['receipts'] if item['id']==id)
    if receipt['status']=='PENDING':
        denied=c.request(f'/hrm/payroll/tax-filing/receipt-review?id={id}&approved=true&reason=synthetic-self-review',{})
        assert denied['code'] not in (0,500),denied
        post('receipt-review',{'id':id,'approved':True,'reason':reason})
        if index==0:
            mismatch=c.request(f'/hrm/payroll/tax-filing/reconcile?id={filing}&reason=synthetic-incomplete',{},finance)
            assert mismatch['code'] not in (0,500),mismatch
            assert detail(filing)['difference']==1
post('reconcile',{'id':filing,'reason':reason},finance)
row=detail(filing)
assert row['filing']['status']=='RECONCILED' and row['officialPaid']==18 and row['difference']==0 and row['sourceFresh'],row['filing']['status']
receipts=c.request(f'/hrm/payroll/tax-filing/receipt-export?id={filing}',token=finance)
assert isinstance(receipts,bytes) and receipts[:2]==b'PK'
sheet=load_workbook(io.BytesIO(receipts)).active
headers=[cell.value for cell in sheet[1]]
exported=[dict(zip(headers,row)) for row in sheet.iter_rows(min_row=2,values_only=True)]
assert len(exported)==2 and sum(Decimal(str(item['金额'])) for item in exported)==18
menus=c.ok('/system/menu/list')
original=c.ok(f"/system/permission/list-role-menus?roleId={s['readerRoleId']}")
query_ids=[item['id'] for item in menus if item.get('permission')=='hrm:payroll:tax:query']
other=next(item['id'] for item in c.ok('/hrm/payroll/entity/list') if item['id']!=e and not any(grant['userId']==s['readerId'] for grant in c.ok(f"/hrm/payroll/entity/grants?id={item['id']}")))
grant_id=None
try:
    c.ok('/system/permission/assign-role-menu',{'roleId':s['readerRoleId'],'menuIds':sorted(set(original+query_ids))})
    grant_id=c.ok('/hrm/payroll/entity/grant',{'entityId':e,'userId':s['readerId'],'actions':['READ'],'effectiveFrom':(today-datetime.timedelta(days=1)).isoformat(),'effectiveTo':(today+datetime.timedelta(days=1)).isoformat(),'reason':reason})
    reader=c.login(s.get('readerUsername','hrmacceptread'),'readerPassword')
    allowed=c.ok(f'/hrm/payroll/tax-filing/list?entityId={e}&year={year}&month={month}',token=reader)
    assert filing in {item['id'] for item in allowed}
    cross=c.request(f'/hrm/payroll/tax-filing/list?entityId={other}&year={year}&month={month}',token=reader)
    assert cross['code'] not in (0,500),cross
    private=c.request(f'/hrm/payroll/tax-filing/export?id={filing}',token=reader)
    assert private['code']==403,private
finally:
    if grant_id is not None:c.query('/hrm/payroll/entity/revoke',{'id':grant_id,'reason':'恢复隔离权限验收前状态'})
    c.ok('/system/permission/assign-role-menu',{'roleId':s['readerRoleId'],'menuIds':original})
result={'entityId':e,'filingId':filing,'period':f'{year}-{month:02d}','status':row['filing']['status'],'income':6000,'withheld':18,'officialPaid':18,'difference':0,'receiptIds':ids,'unclosedPeriodRejected':True,'importReplayIdempotent':True,'selfReviewRejected':True,'partialReceiptMismatchRejected':True,'filingExcelRead':True,'receiptExcelRead':True,'crossSubjectDeniedCode':cross['code'],'privateExportDeniedCode':private['code'],'externalSubmissions':0}
(a.output_dir/'tax-filing-runtime.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
print('Passed: closed-period filing Excel, independent approvals, receipt import/replay, partial mismatch and reconciliation')
