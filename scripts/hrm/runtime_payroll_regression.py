#!/usr/bin/env python3
"""Real API payroll acceptance on isolated synthetic fixtures; failures block delivery."""
import argparse
import datetime
import json
import io
from decimal import Decimal
from concurrent.futures import ThreadPoolExecutor
from openpyxl import Workbook, load_workbook
from pathlib import Path
from runtime_client import RuntimeClient

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--state-file', required=True)
parser.add_argument('--output-dir', type=Path, required=True)
parser.add_argument('--phase', choices=['sources', 'payroll', 'slips'], default='sources')
args = parser.parse_args()
c = RuntimeClient(args.state_file)
s = c.state
entity, batch = s['entityId'], s['batchId']
employees = s['scopeEmployeeIds']
year, month = s.get('fixtureYear', 2026), s.get('fixtureMonth', 10)
payment_date = s.get('fixturePaymentDate', '2026-10-17')
from_date, to_date = f'{year}-01-01', f'{year}-12-31'
opening_date = (datetime.date(year, month, 1) - datetime.timedelta(days=1)) if month > 1 else datetime.date(year, 1, 1)
opening_months = month - 1
review = c.login('hrmacceptreview', 'reviewerPassword')
finance = c.login('hrmacceptoperator', 'browserPassword')
evidence = '虚构隔离验收资料，不作为正式业务签字'
args.output_dir.mkdir(parents=True, exist_ok=True)

def remember(key, create):
    if key not in s:
        s[key] = create()
        c.save()
    return s[key]

def approve(path, id):
    c.query(path, {'id': id, ('approve' if path == '/hrm/payroll/input/review' else 'approved'): True, 'reason': evidence}, review)

if args.phase == 'sources':
    if not c.ok('/hrm/salary/config/get'):
        c.ok('/hrm/salary/config/create', {'cycleStartDay': 1, 'socialSecurityMonthType': 1, 'startYear': year, 'startMonth': month})
    rule = remember('e2eTaxRuleId', lambda: c.ok('/hrm/salary/tax-rule/create', {'name': '隔离累计计税规则', 'type': 1, 'taxEnabled': True, 'threshold': 5000, 'decimalScale': 2, 'cycleType': 2}))
    remember('e2eSalaryGroupId', lambda: c.ok('/hrm/salary/group/create', {'name': '隔离全链路薪资组', 'taxRuleId': rule, 'deptIds': [], 'employeeIds': employees}))
    def policy_request():
        request = c.ok(f'/hrm/payroll/tax-policy/defaults?entityId={entity}&year={year}')
        request.update(requestKey='ACCEPT_TAX_POLICY_2026', officialReference=evidence, evidence=evidence)
        return c.ok('/hrm/payroll/tax-policy/create', request)
    policy = remember('e2eTaxPolicyId', policy_request)
    if next(row['status'] for row in c.ok(f'/hrm/payroll/tax-policy/list?entityId={entity}&year={year}') if row['id'] == policy) != 'APPROVED':
        approve('/hrm/payroll/tax-policy/review', policy)
    for employee in employees:
        proposal = remember(f'e2eSalaryProposal{employee}', lambda: c.ok('/hrm/salary/employee-info/update', {'employeeId': employee, 'recordType': 1, 'changeReason': 0, 'effectTime': f'{year}-{month:02d}-01T00:00:00', 'remark': evidence, 'salaryOptions': [{'code': 10101, 'value': 6000}], 'probationSalaryOptions': [{'code': 10101, 'value': 6000}]}, method='PUT'))
        status = next(row['status'] for row in c.ok('/hrm/payroll/salary-proposal/list') if row['id'] == proposal)
        if status == 'HR_REVIEW':
            c.ok('/hrm/payroll/salary-proposal/review', {'id': proposal, 'reason': evidence}, review)
            status = 'FINANCE_REVIEW'
        if status == 'FINANCE_REVIEW':
            c.ok('/hrm/payroll/salary-proposal/approve', {'id': proposal, 'reason': evidence}, finance)
        work = remember(f'e2eWorkRule{employee}', lambda: c.ok('/hrm/payroll/work/rule', {'entityId': entity, 'employeeId': employee, 'systemType': 'STANDARD', 'effectiveFrom': from_date, 'effectiveTo': to_date, 'legalBasis': evidence, 'cycleMinutes': 9600, 'cycleMonths': 1, 'timeOffExpiryDays': 180, 'limitMinutes': 2160, 'hourlyBase': 50, 'workdayRate': 1.5, 'restdayRate': 2, 'holidayRate': 3}))
        if next(row['status'] for row in c.ok(f'/hrm/payroll/work/rules?entityId={entity}&employeeId={employee}') if row['id'] == work) != 'APPROVED':
            c.query('/hrm/payroll/work/rule/approve', {'id': work}, review)
        if not c.ok(f'/hrm/employee/salary-card/get?employeeId={employee}'):
            c.ok('/hrm/employee/salary-card/save', {'employeeId': employee, 'bankCardNumber': f'TESTBANK{employee:08d}', 'bankName': '虚构隔离银行'}, method='PUT')
        profile = remember(f'e2eTaxProfile{employee}', lambda: c.ok('/hrm/payroll/tax-profile/create', {'entityId': entity, 'employeeId': employee, 'year': year, 'requestKey': f'ACCEPT_PROFILE_{employee}', 'taxpayerName': next(row['name'] for row in c.ok('/hrm/payroll/salary-proposal/employees') if row['id'] == employee), 'identifierType': 'PASSPORT', 'identifierCountry': 'CN', 'identifier': f'TESTONLY{employee:08d}', 'residentThisYear': True, 'noPriorContinuousRemunerationBeforeEntry': False, 'firstSalaryThisYear': False, 'noPriorSalaryThisYear': False, 'separateBonusUsedElsewhere': False, 'openingMode': 'DECLARE', 'openingAsOfDate': opening_date.isoformat(), 'openingIncome': 6000 * opening_months, 'openingSocialDeduction': 400 * opening_months, 'openingBasicDeduction': 5000 * opening_months, 'openingWithheld': 18 * opening_months, 'officialReference': evidence, 'evidence': evidence}))
        if next(row['status'] for row in c.ok(f'/hrm/payroll/tax-profile/list?entityId={entity}&year={year}&employeeId={employee}') if row['id'] == profile) != 'APPROVED':
            approve('/hrm/payroll/tax-profile/review', profile)
        deduction = remember(f'e2eDeduction{employee}', lambda: c.ok('/hrm/payroll/tax-deduction/create', {'entityId': entity, 'employeeId': employee, 'year': year, 'month': month, 'requestKey': f'ACCEPT_DEDUCTION_{employee}', 'sourceDate': f'{year}-{month:02d}-01', 'mode': 'ZERO', 'values': {str(code): 0 for code in range(260101,260107)}, 'officialReference': evidence, 'evidence': evidence}))
        if next(row['status'] for row in c.ok(f'/hrm/payroll/tax-deduction/list?entityId={entity}&year={year}&month={month}&employeeId={employee}') if row['id'] == deduction) != 'APPROVED':
            approve('/hrm/payroll/tax-deduction/review', deduction)
        print(f'Approved salary, work and tax sources for synthetic employee {employee}', flush=True)
    c.ok('/hrm/payroll/batch/payment-plan', {'batchId': batch, 'expectedPaymentDate': payment_date, 'reason': evidence})
    for employee in employees:
        for kind, mode, values in [
            ('ATTENDANCE', 'IMPORT', dict(ACTUAL_DAYS=22, NEED_DAYS=22, **{str(code): 0 for code in range(190101,190107)})),
            ('OVERTIME', 'ZERO', {}), ('DEDUCTIONS', 'SYNC', {}), ('CUMULATIVE_TAX', 'CASH', {})]:
            key = f'e2eInput{employee}{kind}'
            id = remember(key, lambda: c.ok('/hrm/payroll/input/submit', {'batchId': batch, 'employeeId': employee, 'kind': kind, 'mode': mode, 'values': values, 'evidence': evidence}))
            if next(row['status'] for row in c.ok(f'/hrm/payroll/input/list?batchId={batch}') if row['id'] == id) != 'APPROVED':
                approve('/hrm/payroll/input/review', id)
        print(f'Approved all four payroll input kinds for synthetic employee {employee}', flush=True)
    readiness = c.ok(f'/hrm/payroll/input/readiness?batchId={batch}')
    assert len(readiness) == len(employees) * 4 and all(row['status'] == 'APPROVED' for row in readiness), readiness
    (args.output_dir / 'f08-source-readiness.json').write_text(json.dumps(readiness, ensure_ascii=False, indent=2))
    print('Passed: real source approvals and input readiness', flush=True)
elif args.phase == 'payroll':
    current = c.ok(f'/hrm/salary/month-record/get?id={batch}')
    if current['status'] == 5:
        c.query('/hrm/salary/month-record/compute', {'id': batch})
    versions = c.ok(f'/hrm/payroll/batch/versions?batchId={batch}')
    run = versions[0]
    s['e2eRunId'] = run['id']
    c.save()
    rows = json.loads(run['resultSnapshot'])
    assert {row['employeeId'] for row in rows} == set(employees), rows
    for row in rows:
        assert row['expectedPaySalary'] == 6000, row['expectedPaySalary']
        assert row['personalTax'] == 18, row['personalTax']
        assert row['realPaySalary'] == 5582, row['realPaySalary']
    print('Passed: both payroll results match independent gross / tax / net expectations', flush=True)
    for status, action, token in [(11, 'submit', c.token), (12, 'review', review), (13, 'approve', finance), (14, 'freeze', c.token)]:
        current = c.ok(f'/hrm/salary/month-record/get?id={batch}')
        if current['status'] == status:
            c.ok('/hrm/payroll/batch/transition/' + action, {'batchId': batch, 'runId': run['id'], 'reason': evidence}, token)
    current = c.ok(f'/hrm/salary/month-record/get?id={batch}')
    assert current['status'] in (15,16,17,10), current['status']
    (args.output_dir / 'f08-payroll-runtime.json').write_text(json.dumps({'batchId': batch, 'runId': run['id'], 'status': current['status'], 'employees': [{k: row[k] for k in ('employeeId', 'expectedPaySalary', 'personalTax', 'realPaySalary')} for row in rows]}, ensure_ascii=False, indent=2))
    print('Passed: independent HR / finance review and source freeze', flush=True)

    before_versions = [row['id'] for row in versions]
    refused = c.request(f'/hrm/salary/month-record/compute?id={batch}', {})
    assert refused['code'] not in (0,500), refused
    assert [row['id'] for row in c.ok(f'/hrm/payroll/batch/versions?batchId={batch}')] == before_versions
    payments = c.ok(f'/hrm/payroll/payment/list?batchId={batch}')
    if not payments:
        c.query('/hrm/payroll/payment/prepare', {'batchId': batch, 'runId': run['id']}, finance)
        payments = c.ok(f'/hrm/payroll/payment/list?batchId={batch}')
    assert len(payments) == 2
    accounts = {}
    if any(row['status'] == 'PENDING' for row in payments):
        exported = c.request(f'/hrm/payroll/payment/export?batchId={batch}', token=finance)
        assert isinstance(exported, bytes) and exported[:2] == b'PK'
        sheet = load_workbook(io.BytesIO(exported)).active
        headers = [cell.value for cell in sheet[1]]
        for values in sheet.iter_rows(min_row=2, values_only=True):
            item = dict(zip(headers, values))
            accounts[item['代发明细编号']] = item['银行账户']
    def receipt(rows):
        book = Workbook()
        sheet = book.active
        sheet.append(['代发明细编号','发放次数','员工姓名','银行','银行账户','金额','回盘状态','失败原因','实际到账日期'])
        for row, status in rows:
            sheet.append([row['id'],row['attempt'],row['employeeName'],row['bankName'],accounts[row['id']],row['amount'],status, '虚构账号失败' if status == 'FAILED' else '', payment_date if status == 'SUCCESS' else None])
        output = io.BytesIO()
        book.save(output)
        return output.getvalue()
    def upload(action, content):
        result = c.upload(f'/hrm/payroll/payment/{action}?batchId={batch}', 'isolated-receipt.xlsx', content, finance)
        assert result['code'] == 0, result
        return result['data']
    if all(row['status'] == 'PENDING' and row['attempt'] == 1 for row in payments):
        exported = c.request(f'/hrm/payroll/payment/export?batchId={batch}', token=finance)
        assert isinstance(exported, bytes) and exported[:2] == b'PK'
        book = load_workbook(io.BytesIO(exported))
        assert book.active.max_row == 3
        initial = receipt([(payments[0], 'SUCCESS'), (payments[1], 'FAILED')])
        differences = upload('validate', initial)
        assert differences == [], differences
        wrong = dict(payments[0], amount=payments[0]['amount'] + 1)
        invalid = receipt([(wrong, 'SUCCESS')])
        assert upload('validate', invalid), 'Amount mismatch must produce a difference'
        rejected = c.upload(f'/hrm/payroll/payment/reconcile?batchId={batch}', 'invalid-amount.xlsx', invalid, finance)
        assert rejected['code'] not in (0,500), rejected
        with ThreadPoolExecutor(max_workers=2) as executor:
            list(executor.map(lambda _: upload('reconcile', initial), range(2)))
        upload('reconcile', initial)
        assert len(c.ok(f'/hrm/payroll/payment/receipts?batchId={batch}')) == 1
        s['e2eConcurrentReplayPassed'] = True
        s['e2eAmountMismatchRejected'] = True
        c.save()
        print('Passed: real bank Excel export, partial failure and replay idempotency', flush=True)
    payments = c.ok(f'/hrm/payroll/payment/list?batchId={batch}')
    for payment in payments:
        if payment['status'] == 'FAILED':
            c.query('/hrm/payroll/payment/retry', {'id': payment['id']}, finance)
    payments = c.ok(f'/hrm/payroll/payment/list?batchId={batch}')
    pending = [row for row in payments if row['status'] == 'PENDING']
    if pending:
        exported = c.request(f'/hrm/payroll/payment/export?batchId={batch}', token=finance)
        sheet = load_workbook(io.BytesIO(exported)).active
        headers = [cell.value for cell in sheet[1]]
        for values in sheet.iter_rows(min_row=2, values_only=True):
            item = dict(zip(headers, values))
            accounts[item['代发明细编号']] = item['银行账户']
        assert len(pending) == 1 and pending[0]['attempt'] == 2, pending
        stale = dict(pending[0], attempt=1)
        assert upload('validate', receipt([(stale, 'SUCCESS')])), 'Old attempt must be rejected'
        s['e2eStaleAttemptRejected'] = True
        c.save()
        assert upload('validate', receipt([(pending[0], 'SUCCESS')])) == []
        upload('reconcile', receipt([(pending[0], 'SUCCESS')]))
    payments = c.ok(f'/hrm/payroll/payment/list?batchId={batch}')
    assert all(row['status'] == 'SUCCESS' and row['valueDate'] == payment_date for row in payments)
    assert sum(Decimal(str(row['amount'])) for row in payments) == Decimal('11164')
    print('Passed: failed payment retried once and all receipts settled to 11164', flush=True)
    allocation = remember('e2eCostRule', lambda: c.ok('/hrm/payroll/cost/rule', {'entityId':entity, 'employeeId':employees[0], 'effectiveFrom':from_date, 'effectiveTo':to_date, 'employerCost':0, 'parts':[{'deptId':103,'percent':60,'costCenter':'ACCEPT_A'}, {'deptId':104,'percent':40,'costCenter':'ACCEPT_B'}], 'evidence':evidence, 'requestKey':'ACCEPT_COST_60_40'}))
    if next(row['status'] for row in c.ok(f'/hrm/payroll/cost/rules?entityId={entity}') if row['id'] == allocation) != 'APPROVED':
        approve('/hrm/payroll/cost/review-rule', allocation)
    report = remember('e2eCostReport', lambda: c.query('/hrm/payroll/cost/build', {'batchId': batch, 'reason': evidence}))
    reports = c.ok(f'/hrm/payroll/cost/reports?batchId={batch}')
    if next(row['status'] for row in reports if row['id'] == report) != 'APPROVED':
        approve('/hrm/payroll/cost/review-report', report)
    costs = c.ok(f'/hrm/payroll/cost/rows?id={report}')
    assert sum(Decimal(str(row['gross'])) for row in costs) == Decimal('12000')
    assert sum(Decimal(str(row['insurance'])) for row in costs) == Decimal('1600')
    assert sum(Decimal(str(row['total'])) for row in costs) == Decimal('13600')
    assert sorted(Decimal(str(row['total'])) for row in costs if row['employeeId'] == employees[0]) == [Decimal('2720'), Decimal('4080')]
    analysis = c.ok(f'/hrm/payroll/financial/analysis?entityId={entity}&year={year}&month={month}')
    assert analysis['ready'], analysis
    entity_row = next(row for row in analysis['rows'] if row['dimension'] == 'ENTITY')
    assert entity_row['actualCost'] == 13600 and entity_row['budget'] == 100000
    assert entity_row['costRevenueRate'] == 2.72 and entity_row['costPerEmployee'] == 6800, entity_row
    assert entity_row['budgetVariance'] == -86400 and entity_row['budgetVarianceRate'] == -86.4, entity_row
    assert entity_row['costPerPosition'] == 680 and entity_row['staffingVariance'] == -18, entity_row
    print('Passed: approved cost 13600, budget 100000 and revenue ratio 2.72%', flush=True)
    current = c.ok(f'/hrm/salary/month-record/get?id={batch}')
    if current['status'] == 17:
        c.ok('/hrm/payroll/batch/transition/archive', {'batchId': batch, 'runId': run['id'], 'reason': evidence}, finance)
    assert c.ok(f'/hrm/salary/month-record/get?id={batch}')['status'] == 10
    exported_reports = []
    for name, path, column in [
        ('cost', f'/hrm/payroll/cost/export?id={report}', '总成本（元）'),
        ('budget', f'/hrm/payroll/financial/export?entityId={entity}&year={year}&month={month}&dimension=ENTITY', '实际成本（元）')]:
        content = c.request(path)
        assert isinstance(content, bytes) and content[:2] == b'PK', name
        sheet = load_workbook(io.BytesIO(content)).active
        headers = [cell.value for cell in sheet[1]]
        exported_rows = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
        assert sum(Decimal(str(row[column])) for row in exported_rows) == Decimal('13600'), name
        exported_reports.append({'file':name, 'rows':len(exported_rows), 'column':column, 'sum':13600})
    (args.output_dir / 'f08-report-exports.json').write_text(json.dumps({'batchId':batch, 'reports':exported_reports}, ensure_ascii=False, indent=2))
    print('Passed: actual cost and budget Excel files match approved total 13600', flush=True)
    (args.output_dir / 'f08-payment-cost-runtime.json').write_text(json.dumps({'batchId':batch, 'payments':[{k:row[k] for k in ('id','employeeId','amount','status','attempt','valueDate')} for row in payments], 'costRows':costs, 'analysis':analysis, 'concurrentReplayPassed':s.get('e2eConcurrentReplayPassed',False), 'amountMismatchRejected':s.get('e2eAmountMismatchRejected',False), 'staleAttemptRejected':s.get('e2eStaleAttemptRejected',False), 'frozenRecomputeRejected':True, 'archived':c.ok(f'/hrm/salary/month-record/get?id={batch}')['status'] == 10}, ensure_ascii=False, indent=2))

else:
    menus = c.ok('/system/menu/list')
    original = c.ok(f"/system/permission/list-role-menus?roleId={s['readerRoleId']}")
    portal = [row['id'] for row in menus if row.get('permission') == 'hrm:portal:query']
    assert portal
    try:
        c.ok('/system/permission/assign-role-menu', {'roleId':s['readerRoleId'], 'menuIds':sorted(set(original + portal))})
        reader = c.login('hrmacceptread', 'readerPassword')
        def publish():
            return c.ok('/hrm/salary/slip-send-record/create', {'monthRecordId':batch, 'hideEmpty':False, 'all':False, 'employeeIds':[employees[1]], 'channels':['STATION'], 'options':[{'name':'基本工资','type':2,'code':10101,'hidden':False,'sort':1}, {'name':'实发工资','type':2,'code':240101,'hidden':False,'sort':2}]}, finance)
        publication = remember('e2eSlipPublication', publish)
        denied_code = s.get('e2eNonEmployeeSlipDeniedCode')
        if not s.get('e2eSlipWithdrawn'):
            slips = c.ok('/hrm/portal/salary/slip/list', token=reader)
            slip = next(row for row in slips if row['sendRecordId'] == publication)
            detail = c.ok(f"/hrm/portal/salary/slip/get?id={slip['id']}", token=reader)
            assert detail['realPaySalary'] == 5582, detail
            denied = c.request(f"/hrm/portal/salary/slip/get?id={slip['id']}", token=finance)
            assert denied['code'] not in (0,500), denied
            c.ok(f"/hrm/portal/salary/slip/read?ids={slip['id']}", {}, reader, method='PUT')
            if 'e2eSlipWithdrawn' not in s:
                c.ok(f'/hrm/salary/slip-send-record/delete?id={publication}', token=finance, method='DELETE')
                s['e2eSlipWithdrawn'] = True
                c.save()
            denied_after = c.request(f"/hrm/portal/salary/slip/get?id={slip['id']}", token=reader)
            assert denied_after.get('data') is None or denied_after['code'] not in (0,500), denied_after
            assert all(row['id'] != slip['id'] for row in c.ok('/hrm/portal/salary/slip/list', token=reader))
            denied_code = denied['code']
            s['e2eNonEmployeeSlipDeniedCode'] = denied_code
            c.save()
        active = remember('e2eActiveSlipPublication', publish)
        active_slip = next(row for row in c.ok('/hrm/portal/salary/slip/list', token=reader) if row['sendRecordId'] == active)
        assert active_slip['realPaySalary'] == 5582
        s['e2eActiveSlipId'] = active_slip['id']
        c.save()
        (args.output_dir / 'f08-slip-runtime.json').write_text(json.dumps({'employeeId':employees[1], 'activeSlipId':active_slip['id'], 'realPaySalary':active_slip['realPaySalary'], 'otherAccountDeniedCode':denied_code, 'withdrawnSlipHidden':True, 'channels':['STATION'], 'outboundJobsDisabled':True}, ensure_ascii=False, indent=2))
        print('Passed: owner slip details, other account denial, mark-read, withdrawal and republication', flush=True)
    finally:
        c.ok('/system/permission/assign-role-menu', {'roleId':s['readerRoleId'], 'menuIds':original})
